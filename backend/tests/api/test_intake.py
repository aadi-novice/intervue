"""Bulk resume intake, dedupe, cross-role matching."""

import io
import re
import zipfile
from unittest.mock import patch

import pytest
from docx import Document
from moto import mock_aws

from app.core import storage
from app.core.config import settings
from app.models.candidate import Candidate
from app.models.resume_profile import ResumeProfile
from app.models.rubric import RoleMatch
from tests.factories import make_candidate, make_role
from tests.rank_helpers import fake_ai, fake_vec, make_profile


def pdf(text: str) -> bytes:
    return b"%PDF-1.4\n" + text.encode()


def docx(text: str) -> bytes:
    d = Document()
    d.add_paragraph(text)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def zipped(entries: dict[str, bytes]) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        for name, data in entries.items():
            z.writestr(name, data)
    return buf.getvalue()


def fake_extract(data: bytes, filename: str = "") -> str:
    if data.startswith(b"%PDF"):
        return data.decode(errors="ignore").split("\n", 1)[1]
    return "docx " + " ".join(p.text for p in Document(io.BytesIO(data)).paragraphs)


def fake_parse(text: str, org_id: int) -> dict:
    m = re.search(r"EMAIL:(\S+)", text)
    name = re.search(r"NAME:([A-Za-z ]+?)(?: EMAIL|$)", text)
    return {
        "basics": {"full_name": name.group(1).strip() if name else "", "headline": "", "email": m.group(1) if m else "", "phone": ""},
        "summary": text[:80], "education": [], "experience": [], "projects": [],
        "skills": {"programming_languages": [], "frameworks": [], "databases": [], "tools": []},
        "metadata": {"resume_pages": 1, "total_experience_years": 2},
    }


@pytest.fixture(autouse=True)
def _env():
    with mock_aws():
        storage.reset_client()
        storage.ensure_bucket()
        with fake_ai(), patch("app.services.intake_service.extract_text_from_bytes", side_effect=fake_extract), patch(
            "app.services.intake_service.extract_resume_profile", side_effect=fake_parse
        ), patch("app.services.ranking_service.schedule_rank") as sched:
            yield sched
    storage.reset_client()


@pytest.fixture()
def role(db_session, org):
    return make_role(db_session, org.id, "Rust Engineer", description="Rust")


def upload(client, role, files):
    return client.post(f"/roles/{role.id}/resumes/bulk", files=[("files", f) for f in files])


class TestBulkUpload:
    def test_creates_candidates_and_profiles_then_triggers_one_rank(self, client, role, db_session, org, _env):
        r = upload(client, role, [
            ("a.pdf", pdf("NAME:Ann Lee EMAIL:ann@example.com rust"), "application/pdf"),
            ("b.pdf", pdf("NAME:Ben Ray EMAIL:ben@example.com python"), "application/pdf"),
            ("c.docx", docx("NAME:Cy Zed EMAIL:cy@example.com react"), "application/octet-stream"),
        ])
        assert r.status_code == 202
        batch = client.get(f"/batches/{r.json()['batch_id']}").json()
        assert (batch["status"], batch["total"], batch["processed"], batch["failed"]) == ("completed", 3, 3, 0)

        cands = db_session.query(Candidate).filter_by(org_id=org.id).order_by(Candidate.email).all()
        assert [(c.email, c.name, c.source, c.role_id) for c in cands] == [
            ("ann@example.com", "Ann Lee", "bulk", role.id),
            ("ben@example.com", "Ben Ray", "bulk", role.id),
            ("cy@example.com", "Cy Zed", "bulk", role.id),
        ]
        assert all(c.resume_path.startswith(f"{org.id}/resumes/") for c in cands)
        profs = db_session.query(ResumeProfile).filter_by(org_id=org.id).all()
        assert len(profs) == 3 and all(p.embedding is not None and p.raw_text and p.content_hash for p in profs)
        _env.assert_called_once_with(org.id, role.id)  # one debounced re-rank for the whole batch

    def test_existing_candidate_is_updated_not_duplicated(self, client, role, db_session, org):
        old = make_candidate(db_session, org.id, role.id, "ann@example.com", name="Ann Lee",
                             resume_path=storage.put(f"{org.id}/resumes/old.pdf", b"%PDF old", "application/pdf"))
        upload(client, role, [("a.pdf", pdf("NAME:Ann Lee EMAIL:ANN@example.com rust v2"), "application/pdf")])
        db_session.expire_all()
        assert db_session.query(Candidate).filter_by(org_id=org.id).count() == 1
        c = db_session.get(Candidate, old.id)
        assert c.resume_path != f"{org.id}/resumes/old.pdf" and storage.exists(c.resume_path)
        assert not storage.exists(f"{org.id}/resumes/old.pdf")  # replaced blob removed
        assert c.resume_profile.raw_text.endswith("v2")

    def test_identical_resume_is_skipped_as_duplicate(self, client, role, db_session, org):
        same = ("a.pdf", pdf("NAME:Ann Lee EMAIL:ann@example.com rust"), "application/pdf")
        b = upload(client, role, [same, ("copy.pdf", same[1], "application/pdf")]).json()["batch_id"]
        batch = client.get(f"/batches/{b}").json()
        assert db_session.query(Candidate).filter_by(org_id=org.id).count() == 1
        assert any(e["status"] == "duplicate" for e in batch["errors"]) and batch["status"] == "completed"

    def test_zip_is_expanded_and_junk_ignored(self, client, role, db_session, org):
        z = zipped({
            "resumes/a.pdf": pdf("NAME:Ann Lee EMAIL:ann@example.com rust"),
            "resumes/sub/b.docx": docx("NAME:Ben Ray EMAIL:ben@example.com"),
            "notes.txt": b"not a resume",
            "__MACOSX/a.pdf": b"junk",
            "resumes/.hidden.pdf": b"junk",
            "folder/": b"",
        })
        r = upload(client, role, [("batch.zip", z, "application/zip")])
        batch = client.get(f"/batches/{r.json()['batch_id']}").json()
        assert (batch["total"], batch["processed"], batch["failed"], batch["status"]) == (2, 2, 0, "completed")
        assert db_session.query(Candidate).filter_by(org_id=org.id).count() == 2
        assert not any(k.endswith(".zip") for k in _keys(org))  # the archive itself is not retained

    def test_corrupt_inner_files_are_reported_per_file(self, client, role, db_session, org):
        z = zipped({"ok.pdf": pdf("NAME:Ann Lee EMAIL:ann@example.com"), "bad.pdf": b"this is not a pdf"})
        batch = client.get(f"/batches/{upload(client, role, [('x.zip', z, 'application/zip')]).json()['batch_id']}").json()
        assert (batch["total"], batch["processed"], batch["failed"]) == (2, 1, 1)
        assert any(e["file"] == "bad.pdf" and e["status"] == "error" for e in batch["errors"])

    def test_resume_without_text_fails_that_file_only(self, client, role, org):
        r = upload(client, role, [
            ("blank.pdf", pdf("   "), "application/pdf"),
            ("ok.pdf", pdf("NAME:Ann Lee EMAIL:ann@example.com"), "application/pdf"),
        ])
        batch = client.get(f"/batches/{r.json()['batch_id']}").json()
        assert (batch["processed"], batch["failed"], batch["status"]) == (1, 1, "completed")
        assert "No text found" in batch["errors"][0]["detail"]

    def test_missing_email_uses_a_placeholder_and_says_so(self, client, role, db_session, org):
        b = upload(client, role, [("a.pdf", pdf("NAME:Ann Lee only a name"), "application/pdf")]).json()["batch_id"]
        c = db_session.query(Candidate).filter_by(org_id=org.id).one()
        assert c.email.startswith("no-email-") and c.email.endswith("@resume.invalid")
        assert any(e["status"] == "no_email" for e in client.get(f"/batches/{b}").json()["errors"])

    @pytest.mark.parametrize("files,code", [
        ([("a.txt", b"hello", "text/plain")], "unsupported_file"),
        ([("a.pdf", b"", "application/pdf")], "bad_file_size"),
    ])
    def test_request_validation(self, client, role, files, code):
        r = upload(client, role, files)
        assert r.status_code == 422 and r.json()["code"] == code

    def test_over_500_files_rejected_up_front(self, client, role):
        files = [(f"{i}.pdf", pdf(f"x{i}"), "application/pdf") for i in range(501)]
        r = upload(client, role, files)
        assert r.status_code == 422 and r.json()["code"] == "too_many_files"

    def test_zip_bomb_guard(self, client, role, db_session, org):
        with patch("app.services.intake_service.MAX_UNZIPPED_BYTES", 10):
            z = zipped({"a.pdf": pdf("NAME:Ann Lee EMAIL:ann@example.com" + "x" * 100)})
            batch = client.get(f"/batches/{upload(client, role, [('x.zip', z, 'application/zip')]).json()['batch_id']}").json()
        assert db_session.query(Candidate).filter_by(org_id=org.id).count() == 0
        assert (batch["status"], batch["total"], batch["failed"]) == ("completed", 1, 1)
        assert "expands to more than" in batch["errors"][0]["detail"]

    def test_permissions_and_isolation(self, client, other_client, role_client, role, db_session):
        f = [("a.pdf", pdf("NAME:Ann Lee EMAIL:ann@example.com"), "application/pdf")]
        assert upload(role_client("interviewer"), role, f).status_code == 403
        assert upload(other_client, role, f).status_code == 404
        b = upload(client, role, f).json()["batch_id"]
        assert other_client.get(f"/batches/{b}").status_code == 404
        assert client.get("/batches/99999").status_code == 404

    def test_rate_limited(self, client, role):
        f = [("a.pdf", pdf("NAME:Ann Lee EMAIL:ann@example.com"), "application/pdf")]
        with patch.object(settings, "limit_bulk_per_hour", 1):
            assert upload(client, role, f).status_code == 202
            assert upload(client, role, f).status_code == 429


def _keys(org):
    resp = storage._client().list_objects_v2(Bucket=settings.s3_bucket, Prefix=f"{org.id}/")
    return [o["Key"] for o in resp.get("Contents", [])]


class TestCrossRole:
    @pytest.fixture()
    def roles(self, db_session, org):
        rust = make_role(db_session, org.id, "Rust Engineer", description="rust")
        rust.jd_embedding = fake_vec("rust rust systems")
        sales = make_role(db_session, org.id, "Sales Lead", description="sales")
        sales.jd_embedding = fake_vec("sales sales sales")
        closed = make_role(db_session, org.id, "Old Rust Role", description="rust", status="closed")
        closed.jd_embedding = fake_vec("rust rust")
        db_session.commit()
        return rust, sales, closed

    def test_matches_open_roles_above_threshold_only(self, client, db_session, org, roles):
        from app.services import ranking_service

        rust, sales, closed = roles
        cand = make_candidate(db_session, org.id, sales.id, "x@example.com", name="Xavier")
        make_profile(db_session, org.id, cand, "rust rust rust engineer")
        ranking_service.index_candidate(db_session, org.id, cand.id)

        rows = db_session.query(RoleMatch).filter_by(candidate_id=cand.id).all()
        assert {m.role_id for m in rows} == {rust.id}  # sales too dissimilar, closed excluded
        assert rows[0].status == "shortlisted" and rows[0].vector_similarity > 0.9

        out = client.get(f"/candidates/{cand.id}/matches").json()
        assert [(m["role_name"], m["is_primary"], m["strong"]) for m in out] == [("Rust Engineer", False, True)]

    def test_scored_matches_sort_first_and_must_have_failures_are_not_strong(self, client, db_session, org, roles):
        rust, sales, _ = roles
        cand = make_candidate(db_session, org.id, sales.id, "y@example.com")
        db_session.add_all([
            RoleMatch(org_id=org.id, candidate_id=cand.id, role_id=rust.id, status="scored", overall_score=82.0, vector_similarity=0.5),
            RoleMatch(org_id=org.id, candidate_id=cand.id, role_id=sales.id, status="scored", overall_score=90.0,
                      must_have_failed=True, vector_similarity=0.6),
        ])
        db_session.commit()
        out = client.get(f"/candidates/{cand.id}/matches").json()
        assert [m["role_name"] for m in out] == ["Sales Lead", "Rust Engineer"]
        assert [m["strong"] for m in out] == [False, True]
        assert out[0]["is_primary"] is True

    def test_matches_are_org_scoped(self, other_client, db_session, org, roles):
        cand = make_candidate(db_session, org.id, roles[0].id, "z@example.com")
        assert other_client.get(f"/candidates/{cand.id}/matches").status_code == 404
