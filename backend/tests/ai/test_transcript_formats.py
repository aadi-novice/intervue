"""Transcript parsing: txt (header + colon styles), VTT, SRT, DOCX."""

import io

import pytest

from app.ai.transcript_formats import (
    TranscriptParseError,
    format_ms,
    parse_timestamp,
    parse_transcript,
)


def _p(text: str, name: str = "t.txt"):
    return parse_transcript(name, text.encode())


class TestTimestamps:
    @pytest.mark.parametrize(
        "raw,ms",
        [("0:00", 0), ("1:05", 65_000), ("10:23", 623_000), ("01:02:03", 3_723_000),
         ("00:00:01.500", 1_500), ("00:00:01,250", 1_250)],
    )
    def test_parse(self, raw, ms):
        assert parse_timestamp(raw) == ms

    def test_invalid(self):
        assert parse_timestamp("hello") is None

    def test_format(self):
        assert format_ms(65_000) == "1:05" and format_ms(3_723_000) == "1:02:03"


class TestHeaderStyleText:
    """The format the app supported before: `0:00 - Name (Role)` then text lines."""

    def test_multiple_segments_with_times_and_role_hints(self):
        segs = _p(
            "0:00 - Priya Sharma (Interviewer)\nHello and welcome.\n"
            "1:15 - Alex Doe (Candidate)\nHi there.\nThanks for having me.\n"
            "3:30 - Priya Sharma (Interviewer)\nTell me about yourself.\n"
        )
        assert [s.speaker_label for s in segs] == ["Priya Sharma", "Alex Doe", "Priya Sharma"]
        assert [s.kind_hint for s in segs] == ["interviewer", "candidate", "interviewer"]
        assert [s.start_ms for s in segs] == [0, 75_000, 210_000]
        assert segs[1].text == "Hi there. Thanks for having me."
        assert segs[0].end_ms == 75_000  # filled from the next start

    def test_long_timestamps(self):
        segs = _p("10:23 - Interviewer\nQuestion?\n1:05:00 - Candidate\nAnswer.\n")
        assert segs[0].start_ms == 623_000 and segs[1].start_ms == 3_900_000


class TestColonStyleText:
    def test_name_colon_lines(self):
        segs = _p("Priya: Tell me about Rust.\nAlex: I built a compiler.\nIt was fun.\n")
        assert [(s.speaker_label, s.text) for s in segs] == [
            ("Priya", "Tell me about Rust."),
            ("Alex", "I built a compiler. It was fun."),
        ]

    def test_bracketed_timestamps(self):
        segs = _p("[00:01:23] Priya: Hello\n[00:01:30] Alex: Hi\n")
        assert [s.start_ms for s in segs] == [83_000, 90_000]
        assert segs[0].text == "Hello"

    def test_unlabelled_text_becomes_unknown_speaker(self):
        segs = _p("just some words here\n")
        assert segs[0].speaker_label == "Unknown"

    def test_empty_raises(self):
        with pytest.raises(TranscriptParseError):
            _p("   \n\n")


class TestCaptions:
    VTT = (
        "WEBVTT\n\n"
        "00:00:01.000 --> 00:00:04.000\n<v Priya Sharma>Tell me about yourself.\n\n"
        "00:00:05.000 --> 00:00:07.500\n<v Alex Doe>I am a backend engineer\n\n"
        "00:00:07.600 --> 00:00:10.000\n<v Alex Doe>with five years of experience.\n"
    )
    SRT = (
        "1\n00:00:01,000 --> 00:00:03,000\nPriya: How was the project?\n\n"
        "2\n00:00:04,000 --> 00:00:06,000\nAlex: It shipped on time.\n"
    )

    def test_vtt_voice_tags_and_merging_of_adjacent_cues(self):
        segs = parse_transcript("a.vtt", self.VTT.encode())
        assert [s.speaker_label for s in segs] == ["Priya Sharma", "Alex Doe"]
        assert segs[1].text == "I am a backend engineer with five years of experience."
        assert (segs[1].start_ms, segs[1].end_ms) == (5_000, 10_000)

    def test_srt_with_speaker_prefix(self):
        segs = parse_transcript("a.srt", self.SRT.encode())
        assert [(s.speaker_label, s.start_ms, s.end_ms) for s in segs] == [
            ("Priya", 1_000, 3_000),
            ("Alex", 4_000, 6_000),
        ]

    def test_vtt_detected_by_content_even_with_txt_extension(self):
        assert len(parse_transcript("x.txt", self.VTT.encode())) == 2

    def test_vtt_html_tags_stripped(self):
        vtt = "WEBVTT\n\n00:00:01.000 --> 00:00:02.000\n<v A>Hello <i>there</i></v>\n"
        assert parse_transcript("a.vtt", vtt.encode())[0].text == "Hello there"


class TestDocx:
    def test_docx_paragraphs(self):
        from docx import Document

        d = Document()
        d.add_paragraph("0:00 - Priya (Interviewer)")
        d.add_paragraph("Welcome.")
        d.add_paragraph("0:10 - Alex (Candidate)")
        d.add_paragraph("Thanks.")
        buf = io.BytesIO()
        d.save(buf)

        segs = parse_transcript("call.docx", buf.getvalue())
        assert [(s.speaker_label, s.text) for s in segs] == [("Priya", "Welcome."), ("Alex", "Thanks.")]

    def test_corrupt_docx_raises(self):
        with pytest.raises(TranscriptParseError):
            parse_transcript("x.docx", b"not a zip")


def test_latin1_fallback():
    segs = parse_transcript("t.txt", "José: olá\n".encode("latin-1"))
    assert segs[0].speaker_label == "José"
