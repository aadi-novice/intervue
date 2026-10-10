"""
Transcript file parsing -> speaker-labelled segments.

Supported inputs: plain text (``0:00 - Name (Role)`` header blocks, ``[hh:mm:ss] Name: text``
or ``Name: text`` lines), WebVTT, SRT and DOCX. Pure functions, no I/O besides reading bytes.
"""

import io
import re
from dataclasses import dataclass

MAX_SEGMENT_CHARS = 4000


@dataclass
class ParsedSegment:
    speaker_label: str
    start_ms: int
    end_ms: int
    text: str
    kind_hint: str = "unknown"  # interviewer | candidate | unknown (from "(Interviewer)" etc.)


class TranscriptParseError(ValueError):
    pass


# ── time helpers ──────────────────────────────────────────────────────────────
_TS = re.compile(r"^(?:(\d{1,2}):)?(\d{1,2}):(\d{2})(?:[.,](\d{1,3}))?$")


def parse_timestamp(value: str) -> int | None:
    """'1:05' -> 65000, '01:02:03.500' -> 3723500, else None."""
    m = _TS.match(value.strip())
    if not m:
        return None
    h, mnt, s, frac = m.groups()
    ms = int((frac or "0").ljust(3, "0")[:3])
    if h is None:  # MM:SS
        return (int(mnt) * 60 + int(s)) * 1000 + ms
    return ((int(h) * 60 + int(mnt)) * 60 + int(s)) * 1000 + ms


def format_ms(ms: int) -> str:
    s = ms // 1000
    h, rem = divmod(s, 3600)
    m, sec = divmod(rem, 60)
    return f"{h}:{m:02d}:{sec:02d}" if h else f"{m}:{sec:02d}"


_ROLE_HINT = re.compile(r"\(([^)]*)\)\s*$")


def _split_speaker(raw: str) -> tuple[str, str]:
    """'Priya Sharma (Interviewer)' -> ('Priya Sharma', 'interviewer')."""
    raw = raw.strip().strip(":").strip()
    hint = "unknown"
    m = _ROLE_HINT.search(raw)
    if m:
        role = m.group(1).strip().lower()
        if "interview" in role or role in {"host", "recruiter", "hr"}:
            hint = "interviewer"
        elif "candidate" in role or "applicant" in role or "interviewee" in role:
            hint = "candidate"
        raw = raw[: m.start()].strip()
    return (raw or "Unknown"), hint


# ── text ──────────────────────────────────────────────────────────────────────
_HEADER = re.compile(r"^\s*\[?(\d{1,2}(?::\d{2}){1,2}(?:[.,]\d{1,3})?)\]?\s*[-–—]\s*(.+?)\s*$")
_BRACKET_LINE = re.compile(
    r"^\s*[\[(]?(\d{1,2}(?::\d{2}){1,2}(?:[.,]\d{1,3})?)[\])]?\s+([^:]{1,60}?):\s*(.*)$"
)
_COLON_LINE = re.compile(r"^\s*([A-Z][\w .'’()-]{0,58}?):\s+(\S.*)$")


def _finalize(segs: list[ParsedSegment]) -> list[ParsedSegment]:
    out: list[ParsedSegment] = []
    for s in segs:
        s.text = re.sub(r"[ \t]+", " ", s.text).strip()
        if not s.text:
            continue
        s.text = s.text[:MAX_SEGMENT_CHARS]
        out.append(s)
    # fill missing end times from the next start
    for i, s in enumerate(out):
        nxt = out[i + 1].start_ms if i + 1 < len(out) else s.start_ms
        if s.end_ms < s.start_ms or s.end_ms == 0:
            s.end_ms = max(nxt, s.start_ms)
    return out


def _parse_text(text: str) -> list[ParsedSegment]:
    lines = text.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    segs: list[ParsedSegment] = []
    cur: ParsedSegment | None = None

    def start(label: str, hint: str, ts: int | None, first: str = "") -> None:
        nonlocal cur
        base = ts if ts is not None else (segs[-1].start_ms if segs else 0)
        cur = ParsedSegment(label, base, 0, first, hint)
        segs.append(cur)

    for line in lines:
        if not line.strip():
            if cur is not None and cur.text:
                cur.text += "\n"
            continue
        m = _HEADER.match(line)
        if m and parse_timestamp(m.group(1)) is not None:
            label, hint = _split_speaker(m.group(2))
            start(label, hint, parse_timestamp(m.group(1)))
            continue
        m = _BRACKET_LINE.match(line)
        if m and parse_timestamp(m.group(1)) is not None:
            label, hint = _split_speaker(m.group(2))
            start(label, hint, parse_timestamp(m.group(1)), m.group(3))
            continue
        m = _COLON_LINE.match(line)
        if m:
            label, hint = _split_speaker(m.group(1))
            if cur is not None and cur.speaker_label == label and not cur.text.strip():
                cur.text = m.group(2)
            else:
                start(label, hint, None, m.group(2))
            continue
        if cur is None:
            start("Unknown", "unknown", 0, line.strip())
        else:
            cur.text += (" " if cur.text and not cur.text.endswith("\n") else "") + line.strip()
    return _finalize(segs)


# ── VTT / SRT ─────────────────────────────────────────────────────────────────
_CUE = re.compile(
    r"^\s*(\d{1,2}:\d{2}(?::\d{2})?[.,]\d{1,3})\s*-->\s*(\d{1,2}:\d{2}(?::\d{2})?[.,]\d{1,3})"
)
_VOICE = re.compile(r"^<v(?:\.[^ >]+)?\s+([^>]+)>(.*)$", re.S)
_TAGS = re.compile(r"</?[^>]+>")


def _parse_cues(text: str) -> list[ParsedSegment]:
    blocks = re.split(r"\n\s*\n", text.replace("\r\n", "\n").replace("\r", "\n").strip())
    segs: list[ParsedSegment] = []
    for block in blocks:
        lines = [ln for ln in block.split("\n") if ln.strip()]
        cue_idx = next((i for i, ln in enumerate(lines) if _CUE.match(ln)), None)
        if cue_idx is None:
            continue
        m = _CUE.match(lines[cue_idx])
        start, end = parse_timestamp(m.group(1)), parse_timestamp(m.group(2))
        body = " ".join(ln.strip() for ln in lines[cue_idx + 1 :])
        label, hint = "Unknown", "unknown"
        v = _VOICE.match(body)
        if v:
            label, hint = _split_speaker(v.group(1))
            body = v.group(2)
        else:
            c = _COLON_LINE.match(body)
            if c:
                label, hint = _split_speaker(c.group(1))
                body = c.group(2)
        body = _TAGS.sub("", body).strip()
        if body:
            segs.append(ParsedSegment(label, start or 0, end or 0, body, hint))
    # WebVTT/SRT often repeat the speaker only on change: carry the last label forward.
    last = "Unknown"
    for s in segs:
        if s.speaker_label == "Unknown":
            s.speaker_label = last
        else:
            last = s.speaker_label
    return _merge_adjacent(_finalize(segs))


def _merge_adjacent(segs: list[ParsedSegment], max_gap_ms: int = 2500) -> list[ParsedSegment]:
    """Join consecutive cues from the same speaker (captions are split mid-sentence)."""
    out: list[ParsedSegment] = []
    for s in segs:
        p = out[-1] if out else None
        if (
            p
            and p.speaker_label == s.speaker_label
            and s.start_ms - p.end_ms <= max_gap_ms
            and len(p.text) + len(s.text) < MAX_SEGMENT_CHARS
        ):
            p.text += " " + s.text
            p.end_ms = max(p.end_ms, s.end_ms)
        else:
            out.append(s)
    return out


# ── DOCX ──────────────────────────────────────────────────────────────────────
def _docx_text(data: bytes) -> str:
    try:
        from docx import Document

        doc = Document(io.BytesIO(data))
    except Exception as e:
        raise TranscriptParseError(f"Could not read DOCX: {e}") from e
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:  # some exports put "Speaker | time | text" in tables
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" ".join(cells))
    return "\n".join(parts)


# ── entry point ───────────────────────────────────────────────────────────────
def parse_transcript(filename: str, data: bytes) -> list[ParsedSegment]:
    name = (filename or "").lower()
    if name.endswith(".docx"):
        text = _docx_text(data)
        segs = _parse_text(text)
    else:
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = data.decode("latin-1")
        if name.endswith(".vtt") or text.lstrip().startswith("WEBVTT") or (
            "-->" in text and name.endswith(".srt")
        ):
            segs = _parse_cues(text)
        elif "-->" in text and _CUE.search(text):
            segs = _parse_cues(text)
        else:
            segs = _parse_text(text)
    if not segs:
        raise TranscriptParseError("No speech found in transcript")
    return segs
